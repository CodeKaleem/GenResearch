"""Assemble verified proposal sections into a DOCX and human checklist."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches

from models.schemas import ProposalOutput


HUMAN_INPUT_TRIGGERS = (
    "primary data",
    "interview",
    "survey",
    "irb",
    "ethics approval",
    "human participants",
)


def human_input_checklist(output: ProposalOutput) -> list[str]:
    """Identify methodology decisions that cannot be safely invented."""
    checklist = list(output.human_input_needed)
    methodology = next((section.text.lower() for section in output.sections if section.section_name == "methodology"), "")
    for trigger in HUMAN_INPUT_TRIGGERS:
        if trigger in methodology:
            item = f"Confirm the methodology detail involving {trigger}."
            if item not in checklist:
                checklist.append(item)
    if any(result.verdict != "supported" for section in output.sections for result in section.verification):
        checklist.append("Review claims marked partial or unsupported by evidence verification.")
    return checklist


def assemble_docx(output: ProposalOutput, path: str | Path) -> str:
    """Write proposal sections, figures, references, and checklist to DOCX."""
    document = Document()
    document.add_heading(output.topic, level=0)
    for section in output.sections:
        document.add_heading(section.section_name.replace("_", " ").title(), level=1)
        for paragraph in re.split(r"\n\s*\n", section.text.strip()):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    for figure in output.figures:
        if Path(figure).exists():
            document.add_picture(figure, width=Inches(5.5))
    document.add_heading("References", level=1)
    for reference in output.references:
        document.add_paragraph(reference, style="List Bullet")
    document.add_heading("Items Requiring Human Input", level=1)
    for item in human_input_checklist(output) or ["No additional manual input was detected."]:
        document.add_paragraph(item, style="List Bullet")
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return str(destination)